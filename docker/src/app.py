import time
import random
import os
import logging
import docx
import csv

from nicegui import ui
from VShowDB import vShowDB

from cupspage import CupsPage
from judgespage import JudgesPage
from villagerspage import VillagersPage
from adjustentrypage import AdjustEntryPage
from local_file_picker import local_file_picker
from year_picker import year_picker
from are_you_sure import are_you_sure
from reprint_stickers_dialog import ReprintStickersDialog

logFormatter = logging.Formatter("%(asctime)s [%(levelname)-5.5s] {%(name)s} %(message)s")
consoleHandler = logging.StreamHandler()
consoleHandler.setFormatter(logFormatter)
logging.basicConfig(
    level=logging.INFO,
    handlers=[consoleHandler,]
    )
    
LOGGER = logging.getLogger(__name__)

def update_values_dbValue(db, values) -> str:
    db.add_new_row(random.randint(1,100000))
    v = db.get_last_row()
    values['dbValue'] = f'{v}'
    LOGGER.info(f'Updated values: {v}')

async def import_classes(db, pages):
    reason = ( "Import Classes\n"
               "This process will read the selected file and use it to replace all class info for this year.\n"
               "This should only be done on a clean (i.e. after CLEAR YEAR has been used)\n"
               "\nPlease confirm you want to continue ..."
              )
    x = await are_you_sure(reason)
    if not x:
        ui.notify('IMPORT CLASSES - cancelled')
        return
    result = await local_file_picker('/data/in', multiple=False)
    file = os.path.join('/data/in',result[0])
    ui.notify(f'Importing from {file}.')
    try:
        EXP_HEADERS = 'Class,Name,Description,Kids,Judge'
        classes = []
        with open(file, encoding='utf-8-sig') as fh:
            csv_reader = csv.reader(fh, dialect='excel',)
            # for c, line in enumerate(x.rstrip() for x in fh):
            #for c, line in enumerate(csv_reader):
            #    cells = line.split(',')
            for c, cells in enumerate(csv_reader):
                if c == 0: # header
                    init_headers = ','.join(cells[0:5])
                    if init_headers != EXP_HEADERS:
                        raise Exception(f'Failed to import {file}, incorrect header, expected {EXP_HEADERS},*, got {",".join(cells)}')
                else:
                    _class = {
                            'class': cells[0],
                            'name': cells[1],
                            'description': cells[2],
                            'kids': cells[3],
                            'judge': cells[4],
                            'cups': list(x for x in cells[5:] if x)
                            }
                    classes.append(_class)
        db.import_classes(classes)
        ui.notify('Import successful.')
        refresh_pages(pages)
    except Exception as x:
        LOGGER.exception(x)
        ui.notify(x)

def refresh_pages(pages):
    for page in pages:
        page.__reset_data__()

async def really(reason):
    return await are_you_sure(reason)

async def select_year(db, year_label, pages):
    year = await year_picker(db)
    if year:
       db.set_year(year)
       year_label.text = year
       refresh_pages(pages)
       ui.notify(f'Year has been set to {year}.')

async def reset_year(db, year, pages):
    reason = ( "Clear Year\n"
               "This will remove all data related to this year: Entrants, Entries, Classes, Results.\n"
               "This should only be done if a completely clean slate is required for this year!\n"
               "\nPlease confirm you want to continue ..."
              )
    x = await are_you_sure(reason)
    if not x:
        ui.notify('CLEAR YEAR - cancelled')
        return
    db.reset_year()
    refresh_pages(pages)
    ui.notify(f'Year {year} has been cleared.')

async def export_friday_data(db):

    def write_cup_section(np_cups, judge_code, doc):
        if judge_code in np_cups:
            doc.add_heading(f'Cup Section', 3)
            doc.add_paragraph('Please also select an entry to be awarded each of the following cups')
            for cup, classes in np_cups[judge_code].items(): 
                doc.add_heading(f'{cup}', 4)
                doc.add_paragraph(f'Selected from:\n\t' + '\n\t'.join(classes))
                doc.add_paragraph('Entry Ref:')
                doc.add_paragraph('Description:')
    def write_class_list(doc, classlist, ):
        doc.add_heading('Summary of your entries',3)
        doc.add_paragraph('\n\t' + '\n\t'.join(classlist))

    try:
        data = db.get_judge_pages_entry()
        np_cups = {}
        for judge_code, name, cup, display_order, classname in data:
            if judge_code not in np_cups: 
                np_cups[judge_code] = {}
            if cup not in np_cups[judge_code]: 
                np_cups[judge_code][cup] = []
            np_cups[judge_code][cup].append(f'{display_order}. {classname}')

        data = db.get_judge_pages_points()

        doc = docx.Document()
        old_judge_code = ''
        for judge_code, name, display_order, classname, c in data:
            if judge_code != old_judge_code:
                if old_judge_code:
                    write_cup_section(np_cups, old_judge_code, doc)
                    doc.add_page_break()
                old_judge_code = judge_code
                doc.add_heading(f'Judging for section: {judge_code}',2)
                if name:
                    doc.add_paragraph(f'Judge: {name}')
                doc.add_paragraph('Please enter the winning entry ref next to the classes')
                doc.add_paragraph()
            doc.add_paragraph(f'{display_order}.\t{classname}\n(There should be {c} entr{"y" if c == 1 else "ies"})')
            doc.add_paragraph('\t1st:\t\t\t2nd:\t\t\t3rd')
            doc.add_paragraph()

        write_cup_section(np_cups, judge_code, doc)

        doc.save('/data/out/judges.docx')
        ui.notify(f'Wrote Judges guide (judges.docx)')

        data = db.get_entry_hashes()
        with open('/data/out/stickers.csv', 'w') as sfh, \
             open('/data/out/labels.csv', 'w') as lfh:

            lfh.write('Name\n')
            sfh.write('ClassNum,ClassName,EntryRef\n')
            doc = docx.Document()
            old_name = None
            classlist = []
            for name, display_order, classname, ref, _ in data:
                if name != old_name:
                    if old_name:
                        write_class_list(doc,classlist)
                        doc.add_page_break()
                        classlist = []
                    old_name = name

                    lfh.write(f'{name}\n')
                    doc.add_heading('Entrant Guide', 2)
                    doc.add_paragraph(f'Welcome {name}, to the Eynsham Village Show.')
                    doc.add_paragraph('Please follow the following instructions:')
                    doc.add_paragraph( 
                      ( 
                         f'{chr(9679)}\tLabel your entry with the sticker provided, and place on the matching table location.\n'
                         "\t(Note that Children's, and Arts and Crafts are located next door in the Scout Hut).\n"
                         f'{chr(9679)}\tDepart the buildings by 10:30, judging commences at 11:00.\n'
                         f'{chr(9679)}\tThe show reopens to the public at 14:00, with prize giving and raffle at 16:00.\n'
                         f'{chr(9679)}\tPlease collect your entries, and hard-won certificates, before 17:00.\n'
                         f'{chr(9679)}\tNote that any entries not collected before 17:00 will be disposed of.\n'
                         f'{chr(9679)}\tAny issues, please talk to one of the friendly committee members.'
                       ) )

                sfh.write(f'{display_order},"{classname}",{ref}\n')
                classlist.append(f'{display_order}.\t{classname} (Entry Ref: {ref})')

        write_class_list(doc,classlist)
        doc.save('/data/out/entrants.docx')

        ui.notify(f'Wrote labels data (labels.csv)')
        ui.notify(f'Wrote stickers data (stickers.csv)')
        ui.notify(f'Wrote Entrants guide (entrants.docx)')

        data = db.get_category_counts()
        with open('/data/out/category_counts.csv', 'w') as fh:
            fh.write('ClassNum,ClassName,Count\n')
            for display_order, classname, c in data:
                fh.write(f'{display_order},"{classname}",{c}\n')

    except Exception as x:
        LOGGER.exception(x)
        ui.notify(f'Problem writing files - {x}')

async def reprint_stickers(db):
    try:
        available = db.get_villagers_with_entries()
        if not available:
            ui.notify('No villagers with entries found')
            return

        dialog = ReprintStickersDialog(available)
        selected = await dialog

        if not selected:
            ui.notify('No villagers selected for reprint')
            return

        # Get all entry hashes and filter to selected villagers
        data = db.get_entry_hashes()

        with open(f'/data/out/stickers.csv', 'w') as sfh, \
             open(f'/data/out/labels.csv', 'w') as lfh:

            lfh.write('Name\n')
            sfh.write('ClassNum,ClassName,EntryRef\n')

            written_names = set()
            for name, display_order, classname, ref, _ in data:
                if name in selected:
                    if name not in written_names:
                        lfh.write(f'{name}\n')
                        written_names.add(name)
                    sfh.write(f'{display_order},"{classname}",{ref}\n')

        ui.notify(f'Wrote stickers.csv and labels.csv')

    except Exception as x:
        LOGGER.exception(x)
        ui.notify(f'Problem with reprint - {x}')

async def export_saturday_data(db):
    def nice(s: str):
        return s if s else ''

    try:

        # Data for Cups
        data = db.get_cup_winners_pts()
        data2 = db.get_cup_winners_best()

        cups = {}
        points = [3,2,1]
        for cup_number, cup, description, display_order, category, is_child, g, s, b in data:
            if cup not in cups:
                cups[cup] = {'cup_number': cup_number, 
                             'description': description, 
                             'type': 'points',
                             'categories': {},
                             'villagers': {} }
            cups[cup]['categories'][category] = (display_order,(g,s,b))
            for c, villager in enumerate((g,s,b)):
                if villager not in cups[cup]['villagers']:
                    cups[cup]['villagers'][villager] = points[c]
                else: 
                    cups[cup]['villagers'][villager] += points[c]

        for cup_number, cup, description, category, villager, entry_description in data2:
            if cup not in cups:
                cups[cup] = {'cup_number': cup_number, 
                             'description': description, 
                             'type': 'pick',
                             'category': category,
                             'villager': villager,
                             'entry_description': entry_description }

        with open('/data/out/cup_data_raw.csv', 'w') as fh:
            fh.write(f'Num,Cup,Class,Gold,Silver,Bronze,Best\n')
            for cup, v in sorted(cups.items(), key=lambda x: x[1]['cup_number']):
                num = v['cup_number']
                if v['type'] == 'pick':
                    category = v['category']
                    villager = v['villager']
                    fh.write(f'{num},{cup},"{category}",,,,{nice(villager)}\n')
                else: #points
                    for category, (_, (g,s,b)) in sorted(v['categories'].items(), key=lambda x:x[1][0]):
                        fh.write(f'{num},{cup},"{category}",{nice(g)},{nice(s)},{nice(b)},\n')
        ui.notify('Exported raw data (cup_data_raw.csv).')

        # cup point summary
        with open('/data/out/cup_data_pts.csv', 'w') as fh:
            fh.write(f'Num,Cup,Name,Score\n')
            for cup, v in sorted(cups.items(), key=lambda x: x[1]['cup_number']):
                if v['type'] == 'pick': continue
                num = v['cup_number']
                for villager, score in sorted(v['villagers'].items(), key=lambda x: x[1], reverse=True):
                    if villager:
                        fh.write(f'{num},{cup},{villager},{score}\n')
        ui.notify('Exported cup score data (cup_data_pts.csv).')

        doc = docx.Document()
        doc.add_heading('Cup Results', 2)
        for cup, v in sorted(cups.items(), key=lambda x: x[1]['cup_number']):
            LOGGER.info(cup)
            num = v['cup_number']
            description = v['description']
            if v['type'] == 'pick': 
                category = v['category']
                villager = v['villager']
                entry = v['entry_description']
                doc.add_paragraph(f'{num}. {cup} ({description})\n{villager} (Awarded for: {nice(entry)})')
            else:
                for villager, score in sorted(v['villagers'].items(), key=lambda x: x[1], reverse=True):
                    if villager: # There is a None villager who gets all the rest of the points
                        doc.add_paragraph(f'{num}. {cup} ({description})\n{villager}')
                        break

        doc.save('/data/out/cup_data.docx')

        ui.notify('Exported cup winners (cup_data.docx).')

        data = db.get_class_winners()
        with open('/data/out/gold.csv', 'w') as gfh, \
             open('/data/out/silver.csv', 'w') as sfh, \
             open('/data/out/bronze.csv', 'w') as bfh:
            gfh.write('ClassNum,ClassName,Name\n')
            sfh.write('ClassNum,ClassName,Name\n')
            bfh.write('ClassNum,ClassName,Name\n')
            for c, name, g, s, b in data:
                if g: gfh.write(f'{c},"{name}",{g}\n')
                if s: sfh.write(f'{c},"{name}",{s}\n')
                if b: bfh.write(f'{c},"{name}",{b}\n')
        ui.notify('Exported certificate data (<medal>.csv).')

    except Exception as x:
        LOGGER.exception(x)
        ui.notify(f'Problem writing files - {x}')


def display_tabs(db):
    pages = []
    with ui.header().classes(replace='row items-center') as header:
        ui.button(on_click=lambda: left_drawer.toggle(), icon='menu').props('flat color=white')
        with ui.tabs() as tabs:
            ui.tab('Welcome')
            ui.tab('Cups')
            ui.tab('Villagers')
            ui.tab('Judges')
            ui.tab('Adjust Entry')
    
    with ui.footer(value=False) as footer:
        ui.label('Just call John')
    
    with ui.left_drawer().classes('bg-blue-100') as left_drawer:
        with ui.column():
            ui.image('resources/logo.png')
            year_label = ui.label(db.get_year())
            ui.button('Set Year', on_click=lambda: select_year(db, year_label, pages))
            ui.button('Clear Year', on_click=lambda: reset_year(db, year_label.text, pages))
            ui.button('Reset Pages', on_click=lambda: refresh_pages(pages))
            ui.button('Import classes', on_click=lambda: import_classes(db, pages))
            ui.button('Export Pre-Show files', on_click=lambda: export_friday_data(db))
            ui.button('Re-export Selected Villagers', on_click=lambda: reprint_stickers(db))
            ui.button('Export Post-Show files', on_click=lambda: export_saturday_data(db))

    with ui.page_sticky(position='bottom-right', x_offset=20, y_offset=20):
        ui.button(on_click=footer.toggle, icon='contact_support').props('fab')
    
    with ui.tab_panels(tabs, value='Welcome').classes('w-full'):
        with ui.tab_panel('Welcome'):
            with open('resources/welcome.md',) as fh:
                ui.markdown(fh.read(), extras=['tables'])

        with ui.tab_panel('Cups'):
            pages.append(CupsPage(db))
        with ui.tab_panel('Villagers'):
            pages.append(VillagersPage(db))
        with ui.tab_panel('Judges'):
            pages.append(JudgesPage(db))
        with ui.tab_panel('Adjust Entry'):
            pages.append(AdjustEntryPage(db))

def db_test(db):
    ui.label('Village Show')
    values = {'dbValue': '0'}
    ui.button('Click To Update', on_click=lambda: update_values_dbValue(db, values))
    ui.button('Click To Test', on_click=lambda: ui.notify('OOh, you pressed me'))
    number = ui.label('0').bind_text_from(values, 'dbValue')

if __name__ in {"__main__", "__mp_main__"}:
    LOGGER.info('Application starting')
    db = vShowDB(scriptdir='/app/sql', run_setup=(__name__=="__main__"))
 
    # db_test(db)
    display_tabs(db)
    
    LOGGER.info('Starting the web server')
    ui.run(port=8080)
    
    #database testing
    #while True:
    #    add_new_row(random.randint(1,100000))
    #    print('The last value insterted is: {}'.format(get_last_row()))
    #    time.sleep(5)
